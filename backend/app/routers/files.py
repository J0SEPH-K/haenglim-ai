import os
import uuid
from fastapi import APIRouter, Depends, UploadFile, File, HTTPException
from fastapi.responses import FileResponse
from app.config import get_settings
from app.models.models import User
from app.utils.dependencies import get_current_user
from app.utils.document import DOCUMENT_EXTENSIONS

router = APIRouter()

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
ALLOWED_EXTENSIONS = IMAGE_EXTENSIONS | DOCUMENT_EXTENSIONS
MAX_FILE_SIZE = 20 * 1024 * 1024  # 20 MB
THUMB_DIR = "thumbnails"


@router.post("/upload")
async def upload_file(
    file: UploadFile = File(...),
    _: User = Depends(get_current_user),
):
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail=f"File type not allowed. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}")

    contents = await file.read()
    if len(contents) > MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail="File too large (max 20MB)")

    original_name = file.filename or "file"
    filename = f"{uuid.uuid4().hex}{ext}"
    settings = get_settings()
    filepath = os.path.join(settings.UPLOAD_DIR, filename)

    with open(filepath, "wb") as f:
        f.write(contents)

    file_type = "document" if ext in DOCUMENT_EXTENSIONS else "image"

    # Generate thumbnail for documents
    thumbnail_url = None
    page_count = None
    if file_type == "document":
        result = _generate_thumbnail(filepath, filename, settings.UPLOAD_DIR)
        if result:
            thumbnail_url = result["thumbnail_url"]
            page_count = result.get("page_count")

    return {
        "url": f"/uploads/{filename}",
        "filename": filename,
        "original_name": original_name,
        "file_type": file_type,
        "thumbnail_url": thumbnail_url,
        "page_count": page_count,
    }


@router.get("/thumbnail/{filename}")
async def get_thumbnail(filename: str):
    settings = get_settings()
    thumb_dir = os.path.join(settings.UPLOAD_DIR, THUMB_DIR)
    # filename can be "abc.png" or "abc_p2.png" etc.
    thumb_path = os.path.join(thumb_dir, filename)
    if not os.path.exists(thumb_path):
        raise HTTPException(status_code=404, detail="Thumbnail not found")
    return FileResponse(thumb_path, media_type="image/png")


def _generate_thumbnail(filepath: str, filename: str, upload_dir: str) -> dict | None:
    """Generate thumbnail images for a document. Returns dict with thumbnail_url and page_count, or None."""
    thumb_dir = os.path.join(upload_dir, THUMB_DIR)
    os.makedirs(thumb_dir, exist_ok=True)
    base_name = os.path.splitext(filename)[0]

    ext = os.path.splitext(filepath)[1].lower()

    try:
        if ext == ".pdf":
            return _thumb_pdf(filepath, thumb_dir, base_name)
        elif ext == ".pptx":
            thumb_path = os.path.join(thumb_dir, f"{base_name}.png")
            url = _thumb_pptx_url(filepath, thumb_path, base_name)
            return {"thumbnail_url": url, "page_count": None} if url else None
        elif ext == ".docx":
            thumb_path = os.path.join(thumb_dir, f"{base_name}.png")
            url = _thumb_docx_url(filepath, thumb_path, base_name)
            return {"thumbnail_url": url, "page_count": None} if url else None
        elif ext == ".xlsx":
            thumb_path = os.path.join(thumb_dir, f"{base_name}.png")
            url = _thumb_xlsx_url(filepath, thumb_path, base_name)
            return {"thumbnail_url": url, "page_count": None} if url else None
    except Exception:
        pass

    return None


def _thumb_pdf(filepath: str, thumb_dir: str, base_name: str) -> dict | None:
    import fitz
    doc = fitz.open(filepath)
    page_count = len(doc)
    if page_count == 0:
        doc.close()
        return None

    # Generate thumbnails for up to 3 pages
    for i in range(min(page_count, 3)):
        page = doc[i]
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
        suffix = "" if i == 0 else f"_p{i + 1}"
        pix.save(os.path.join(thumb_dir, f"{base_name}{suffix}.png"))

    doc.close()
    return {
        "thumbnail_url": f"/api/files/thumbnail/{base_name}.png",
        "page_count": page_count,
    }


def _thumb_pptx_url(filepath: str, thumb_path: str, base_name: str) -> str | None:
    import zipfile
    with zipfile.ZipFile(filepath) as zf:
        for name in zf.namelist():
            if name.startswith("docProps/thumbnail") or name == "docProps/thumbnail.jpeg":
                data = zf.read(name)
                from PIL import Image
                import io
                img = Image.open(io.BytesIO(data))
                img.save(thumb_path, "PNG")
                return f"/api/files/thumbnail/{base_name}.png"
    return None


def _thumb_docx_url(filepath: str, thumb_path: str, base_name: str) -> str | None:
    import zipfile
    with zipfile.ZipFile(filepath) as zf:
        for name in zf.namelist():
            if "thumbnail" in name.lower():
                data = zf.read(name)
                from PIL import Image
                import io
                img = Image.open(io.BytesIO(data))
                img.save(thumb_path, "PNG")
                return f"/api/files/thumbnail/{base_name}.png"
    return None


def _thumb_xlsx_url(filepath: str, thumb_path: str, base_name: str) -> str | None:
    from openpyxl import load_workbook
    from PIL import Image, ImageDraw

    wb = load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active
    if not ws:
        wb.close()
        return None

    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= 8:
            break
        cells = [str(c)[:12] if c is not None else "" for c in row[:5]]
        rows.append(cells)
    wb.close()

    if not rows:
        return None

    w, h = 300, 200
    img = Image.new("RGB", (w, h), "#2f2f2f")
    draw = ImageDraw.Draw(img)

    y = 10
    for i, row in enumerate(rows):
        x = 10
        bg = "#3a3a3a" if i % 2 == 0 else "#333333"
        draw.rectangle([5, y - 2, w - 5, y + 18], fill=bg)
        for cell in row:
            draw.text((x, y), cell, fill="#cccccc")
            x += 58
        y += 22

    img.save(thumb_path, "PNG")
    return f"/api/files/thumbnail/{base_name}.png"
